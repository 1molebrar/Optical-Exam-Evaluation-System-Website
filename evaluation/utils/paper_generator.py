"""PDF answer-sheet generator — exact Topkapi University exam template."""
import os
from datetime import date, datetime
from io import BytesIO
from reportlab.lib.pagesizes import A4
from reportlab.lib.colors import HexColor, black, white
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
import barcode
from barcode.writer import ImageWriter

from django.conf import settings

# Register Roboto fonts
FONTS_DIR = os.path.join(settings.BASE_DIR, 'evaluation', 'fonts')
pdfmetrics.registerFont(TTFont('Roboto', os.path.join(FONTS_DIR, 'Roboto-Regular.ttf')))
pdfmetrics.registerFont(TTFont('Roboto-Bold', os.path.join(FONTS_DIR, 'Roboto-Bold.ttf')))

from .layout_engine import (
    PAGE_WIDTH, PAGE_HEIGHT, MARKER_INSET, MARKER_SIZE,
    CONTENT_LEFT, CONTENT_RIGHT, CONTENT_WIDTH,
    TOP_BOX_TOP, BARCODE_BOX_WIDTH, BARCODE_BOX_HEIGHT, BARCODE_BOX_X,
    SUBHEADER_TOP, INFO_TABLE_TOP, INFO_COL_WIDTHS, INFO_HEADER_HEIGHT, INFO_ROW_HEIGHT, INFO_TABLE_BOTTOM,
    BUBBLE_SHEET_TOP, BUBBLE_RADIUS, BUBBLE_SPACING_X, BUBBLE_SPACING_Y, BUBBLE_SHEET_BOTTOM,
    SCORE_SHEET_TOP, SCORE_HEADER_HEIGHT_1, SCORE_HEADER_HEIGHT_2, SCORE_ROW_HEIGHT, SCORE_TABLE_BOTTOM,
    get_student_id_bubble_positions_pdf
)

NAVY = HexColor('#000000')  # Keep it black or dark gray like the PDF
GRID_COLOR = HexColor('#000000') # PDF uses black borders
TEXT_COLOR = HexColor('#000000')

def draw_alignment_markers(c, page_width, page_height):
    marker_size = MARKER_SIZE
    inset = MARKER_INSET
    corners = [
        (inset, page_height - inset - marker_size),
        (page_width - inset - marker_size, page_height - inset - marker_size),
        (page_width - inset - marker_size, inset),
        (inset, inset),
    ]
    for cx, cy in corners:
        c.setFillColor(black)
        c.rect(cx, cy, marker_size, marker_size, fill=1, stroke=0)

def draw_university_header(c, exam):
    uni = exam.course.university if (exam and hasattr(exam, 'course') and exam.course and exam.course.university) else None
    
    uni_name = uni.name if uni else "İSTANBUL TOPKAPI ÜNİVERSİTESİ"
    
    # Logo text
    c.setFont('Roboto-Bold', 22)
    c.setFillColor(TEXT_COLOR)
    
    # Try to draw the logo if it exists
    logo_drawn = False
    if uni and uni.logo and uni.logo.path and os.path.exists(uni.logo.path):
        try:
            # Maintain aspect ratio. Logo box is max 80x80
            from PIL import Image as PILImage
            with PILImage.open(uni.logo.path) as img:
                w, h = img.size
                ratio = min(80.0/w, 80.0/h)
                draw_w, draw_h = w * ratio, h * ratio
                # Center the logo vertically within the header box, place it on the left
                c.drawImage(uni.logo.path, CONTENT_LEFT, PAGE_HEIGHT - 120 + (80 - draw_h)/2, width=draw_w, height=draw_h, preserveAspectRatio=True)
                logo_drawn = True
        except Exception:
            pass

    text_start_x = CONTENT_LEFT + 100 if logo_drawn else CONTENT_LEFT + 20

    if not logo_drawn and not uni:
        c.drawString(text_start_x, PAGE_HEIGHT - 65, "İSTANBUL")
        c.setFont('Roboto-Bold', 28)
        c.drawString(text_start_x, PAGE_HEIGHT - 90, "TOPKAPI")
        c.setFont('Roboto', 16)
        c.drawString(text_start_x, PAGE_HEIGHT - 108, "ÜNİVERSİTESİ")
        # T Logo fallback
        c.rect(CONTENT_LEFT - 10, PAGE_HEIGHT - 110, 20, 60, fill=1, stroke=0)
        c.rect(CONTENT_LEFT - 20, PAGE_HEIGHT - 65, 40, 15, fill=1, stroke=0)
    else:
        # Custom university name
        c.setFont('Roboto-Bold', 14)
        c.drawString(text_start_x, PAGE_HEIGHT - 55, uni_name)

    # Removed centered SORU FORMU text to avoid overlap

def generate_barcode_image(data):
    try:
        from barcode import Code128
        from barcode.writer import ImageWriter
        buffer = BytesIO()
        Code128(data, writer=ImageWriter()).write(buffer)
        return buffer
    except Exception as e:
        print(f"Barcode gen error: {e}")
        return None

def draw_qr_and_barcode_block(c, layout, exam_paper):
    qr_path = exam_paper.qr_image.path if exam_paper and exam_paper.qr_image else None
    paper_code = exam_paper.paper_code if exam_paper else "UNKNOWN"
    event_id = exam_paper.exam.id if exam_paper else 0

    c.setStrokeColor(black)
    c.setLineWidth(1)
    c.roundRect(BARCODE_BOX_X, TOP_BOX_TOP - BARCODE_BOX_HEIGHT, BARCODE_BOX_WIDTH, BARCODE_BOX_HEIGHT, 4, fill=0, stroke=1)
    
    # Generate and draw Barcode
    bc_buffer = generate_barcode_image(paper_code)
    if bc_buffer:
        c.drawImage(
            from_buffer(bc_buffer), BARCODE_BOX_X + 10, TOP_BOX_TOP - BARCODE_BOX_HEIGHT + 15,
            width=100, height=25, preserveAspectRatio=False
        )

    # Draw Event ID text
    c.setFont('Roboto', 7)
    c.drawCentredString(BARCODE_BOX_X + 60, TOP_BOX_TOP - BARCODE_BOX_HEIGHT + 5, f"Event ID: {event_id} - Page: 1")

    # Draw QR code
    if qr_path and os.path.exists(qr_path):
        qr_size = 35
        c.drawImage(qr_path, BARCODE_BOX_X + BARCODE_BOX_WIDTH - qr_size - 5, TOP_BOX_TOP - BARCODE_BOX_HEIGHT + 5, width=qr_size, height=qr_size)

def from_buffer(buf):
    from reportlab.lib.utils import ImageReader
    buf.seek(0)
    return ImageReader(buf)

def draw_subheaders(c, exam):
    c.setFont('Roboto', 10)
    
    date_text = "Tarih:"
    if exam and exam.exam_date:
        date_text += f" {exam.exam_date.strftime('%d/%m/%Y')}"
    
    # Sağ tarafa yaslı olarak tarihi çiz (right-aligned)
    c.drawRightString(CONTENT_RIGHT, SUBHEADER_TOP, date_text)
    
    # Dersi ve başlığı birleştir
    course_text = f"{exam.course.code} - {exam.course.name} - Final Sınavı : {exam.title}" if exam else "COURSE CODE - COURSE NAME - EXAM"
    
    # Eğer metin çok uzunsa kırp (Overflow control)
    max_width = CONTENT_WIDTH - c.stringWidth(date_text, 'Roboto', 10) - 20
    while c.stringWidth(course_text, 'Roboto', 10) > max_width and len(course_text) > 10:
        course_text = course_text[:-4] + "..."
        
    # Sol tarafa ders ismini çiz
    c.drawString(CONTENT_LEFT, SUBHEADER_TOP, course_text)

def draw_student_info_table(c):
    c.setStrokeColor(black)
    c.setLineWidth(1)
    
    y = INFO_TABLE_TOP
    c.rect(CONTENT_LEFT, y - INFO_HEADER_HEIGHT - INFO_ROW_HEIGHT, CONTENT_WIDTH, INFO_HEADER_HEIGHT + INFO_ROW_HEIGHT)
    c.line(CONTENT_LEFT, y - INFO_HEADER_HEIGHT, CONTENT_RIGHT, y - INFO_HEADER_HEIGHT)
    
    x = CONTENT_LEFT
    labels = ["Öğrencinin Adı \u2013 Soyadı", "Öğrenci No", "Programı/Bölümü", "İmza"]
    c.setFont('Roboto-Bold', 10)
    for w, label in zip(INFO_COL_WIDTHS, labels):
        if x > CONTENT_LEFT:
            c.line(x, y, x, y - INFO_HEADER_HEIGHT - INFO_ROW_HEIGHT)
        c.drawString(x + 5, y - 14, label)
        x += w

def draw_bubble_sheet(c, layout):
    c.setFont('Roboto-Bold', 8)
    grid_x = layout['student_id_grid_x']
    # Center the title over the actual grid (grid width is 11 * 11.5 = 126.5, half is ~63)
    c.drawCentredString(grid_x + 63, BUBBLE_SHEET_TOP, "STUDENT ID BUBBLE SHEET")
    
    c.setStrokeColor(black)
    c.setLineWidth(0.5)
    
    box_width = layout['bubble_spacing_x']
    box_height = 16
    
    positions = get_student_id_bubble_positions_pdf(layout)
    for col_idx, col_positions in enumerate(positions):
        # Draw OCR box at the top
        bx, by, br = col_positions[0]
        box_y = by + br + 4
        box_x = bx - box_width / 2
        c.rect(box_x, box_y, box_width, box_height)
        
        for val, (bx, by, br) in enumerate(col_positions):
            c.circle(bx, by, br, fill=0, stroke=1)
            # Write the digit inside
            c.setFont('Roboto', 6)
            c.drawCentredString(bx, by - 2, str(val))
            
    # Draw a light grid around the bubbles
    bx, by, br = positions[0][0]
    top_y = by + br + 4 + box_height
    bottom_y = positions[0][-1][1] - br - 3
    
    width = layout['student_id_length'] * layout['bubble_spacing_x']
    height = top_y - bottom_y
    start_y = bottom_y
    
    c.rect(grid_x + 3, start_y, width, height)
    # Draw vertical lines separating columns
    for i in range(1, layout['student_id_length']):
        vx = grid_x + 3 + i * BUBBLE_SPACING_X
        c.line(vx, start_y, vx, start_y + height)

def draw_dynamic_questions(c, exam, layout):
    c.setStrokeColor(black)
    c.setLineWidth(0.5)
    
    omr_centers = layout.get('mcq_bubble_centers', [])
    classical_rects = layout.get('classical_score_rects', [])
    
    questions = list(exam.questions.all().order_by('question_number')) if exam else []
    
    omr_idx = 0
    class_idx = 0
    
    for q in questions:
        q_num = q.question_number
        if q.question_type == 'omr':
            if omr_idx < len(omr_centers):
                centers = omr_centers[omr_idx]
                if centers:
                    cy = centers[0][1]
                    c.setFont('Roboto-Bold', 10)
                    c.drawString(CONTENT_LEFT, cy - 3, f"Q{q_num}:")
                    c.setFont('Roboto', 8)
                    for i, (bx, by) in enumerate(centers):
                        c.circle(bx, by, 6, fill=0, stroke=1)
                        c.drawCentredString(bx, by - 3, chr(65 + i)) # A B C D E
                omr_idx += 1
        elif q.question_type == 'classical':
            if class_idx < len(classical_rects):
                x, y_bottom, box_width, box_height = classical_rects[class_idx]
                c.setFont('Roboto-Bold', 10)
                c.drawString(CONTENT_LEFT, y_bottom + box_height/2 - 3, f"Q{q_num} - Score:")
                
                # Draw fiducial marker (black square)
                marker_size = 12
                marker_x = x - marker_size - 6
                marker_y = y_bottom + box_height/2 - marker_size/2
                c.rect(marker_x, marker_y, marker_size, marker_size, fill=1, stroke=0)
                
                c.setStrokeColor(black)
                c.rect(x, y_bottom, box_width, box_height, fill=0, stroke=1)
                c.setFont('Roboto', 8)
                c.drawString(x + box_width + 10, y_bottom + box_height/2 - 3, f"/ {q.points:g} pts")
                class_idx += 1

def draw_exam_rules(c, exam, layout):
    if not exam or not getattr(exam, 'rules', None):
        return
        
    rules_text = exam.rules.strip()
    if not rules_text:
        return
        
    c.setFont('Roboto-Bold', 10)
    # Sınav kurallarını Bubble Sheet'in eski yeri olan 2. sütun hizasına taşıyalım
    start_x = CONTENT_LEFT + 220
    # Doğrudan Bubble Sheet'in tepe noktası hizasından başlasın
    start_y = BUBBLE_SHEET_TOP
    
    c.drawString(start_x, start_y, "Sınav Kuralları / Exam Rules:")
    
    # Simple word wrap
    text_object = c.beginText(start_x, start_y - 15)
    text_object.setFont('Roboto', 8)
    
    max_width = CONTENT_RIGHT - start_x
    
    for line in rules_text.split('\n'):
        words = line.split(' ')
        current_line = ""
        for word in words:
            if c.stringWidth(current_line + word + " ", 'Roboto', 8) < max_width:
                current_line += word + " "
            else:
                text_object.textLine(current_line)
                current_line = word + " "
        text_object.textLine(current_line)
        
    c.drawText(text_object)

def generate_exam_pdf(exam, exam_paper):
    layout = exam.get_sheet_layout() if exam else None
    if not layout:
        from .layout_engine import compute_sheet_layout
        layout = compute_sheet_layout()
    
    buffer = BytesIO()
    c = canvas.Canvas(buffer, pagesize=A4)
    
    draw_alignment_markers(c, PAGE_WIDTH, PAGE_HEIGHT)
    draw_university_header(c, exam)
    draw_qr_and_barcode_block(c, layout, exam_paper)
    draw_subheaders(c, exam)
    draw_student_info_table(c)
    draw_bubble_sheet(c, layout)
    draw_dynamic_questions(c, exam, layout)
    draw_exam_rules(c, exam, layout)
    
    c.save()
    buffer.seek(0)
    return buffer

def pdf_buffer_to_image(pdf_buffer, dpi=300):
    pdf_buffer.seek(0)
    pdf_bytes = pdf_buffer.read()
    try:
        import fitz
        doc = fitz.open(stream=pdf_bytes, filetype='pdf')
        page = doc[0]
        zoom = dpi / 72.0
        pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
        doc.close()
    except ImportError:
        from PIL import Image
        import io
        from pdf2image import convert_from_bytes
        images = convert_from_bytes(pdf_bytes, dpi=dpi)
        out = io.BytesIO()
        images[0].save(out, format='PNG')
        out.seek(0)
        return out
    
    from PIL import Image
    img = Image.frombytes('RGB', (pix.width, pix.height), pix.samples)
    out = BytesIO()
    img.save(out, format='PNG')
    out.seek(0)
    return out

def generate_exam_image(exam, exam_paper, dpi=300):
    pdf_buffer = generate_exam_pdf(exam, exam_paper)
    return pdf_buffer_to_image(pdf_buffer, dpi=dpi)

def generate_exam_docx(exam, exam_paper):
    import docx
    from docx.shared import Cm
    from io import BytesIO
    
    doc = docx.Document()
    
    sections = doc.sections
    for section in sections:
        section.top_margin = Cm(0)
        section.bottom_margin = Cm(0)
        section.left_margin = Cm(0)
        section.right_margin = Cm(0)
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)

    if exam_paper and exam_paper.image_file:
        try:
            doc.add_picture(exam_paper.image_file.path, width=Cm(21.0), height=Cm(29.7))
        except Exception:
            img_buffer = generate_exam_image(exam, exam_paper, dpi=300)
            doc.add_picture(img_buffer, width=Cm(21.0), height=Cm(29.7))
    else:
        img_buffer = generate_exam_image(exam, exam_paper, dpi=300)
        doc.add_picture(img_buffer, width=Cm(21.0), height=Cm(29.7))
        
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
